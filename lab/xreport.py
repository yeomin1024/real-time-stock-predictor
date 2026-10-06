"""1000% 탐색 결과 묶기 → results_x/1000pct_탐색_결과.xlsx · 자산곡선_1000pct.png"""
import os, sys, json
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(os.path.dirname(HERE), "results_x")
plt.rcParams["font.family"] = "Malgun Gothic"
plt.rcParams["axes.unicode_minus"] = False

KEY = [("X01_utilization", "C0", "C0 기본(현금 안, 58종)"),
       ("X07_combo", "U", "U 신용 없이 최선(+레버리지ETF·2회)"),
       ("X08_robust_X2", "X2", "X2 1000% 목표(U + 신용 2배)"),
       ("X07_combo", "U SP100", "U · 2023년 S&P100(편향 점검)"),
       ("X07_combo", "U L2 SP100", "X2 설정 · 2023년 S&P100(편향 점검)")]

allr = pd.read_csv(os.path.join(RES, "all_x_rounds.csv"), encoding="utf-8-sig")
cols = ["라운드", "변형", "종목수", "ALL_수익률(%)", "ALL_연환산(%)", "ALL_MDD(%)", "ALL_샤프", "IS_수익률(%)", "IS_연환산(%)",
        "IS_MDD(%)", "IS_샤프", "OOS_수익률(%)", "OOS_연환산(%)", "OOS_MDD(%)", "OOS_샤프", "ALL_청산횟수", "ALL_승률(%)",
        "ALL_손익비", "평균투자비중", "최대투자비중", "이자($)", "ALL_SPY수익률(%)", "판정"]
allr = allr[[c for c in cols if c in allr.columns]]

curves = {}
for rnd, name, label in KEY:
    p = os.path.join(RES, rnd, f"{name}_equity.csv")
    e = pd.read_csv(p, encoding="utf-8-sig")
    curves[label] = e.assign(날짜=pd.to_datetime(e["날짜"])).set_index("날짜")["자산($)"]
eq = pd.DataFrame(curves)
spy = pd.read_pickle(os.path.join(HERE, "data", "spy_daily.pkl"))
s = spy.reindex(eq.index).ffill()
prev = spy[spy.index < eq.index[0]].iloc[-1]
eq["SPY 단순보유"] = s / prev * 10000

fig, ax = plt.subplots(figsize=(12, 6))
for c in eq.columns:
    ax.plot(eq.index, eq[c], label=c, lw=2.2 if c.startswith("X2 1000") else 1.3,
            ls="--" if "S&P100" in c else ("-." if c.startswith("SPY") else "-"))
ax.axhline(110000, color="red", lw=0.8, ls=":")
ax.text(pd.Timestamp("2025-01-15"), 116000, "목표 +1000% ($110,000)", color="red", fontsize=9)
ax.axvline(pd.Timestamp("2025-08-01"), color="gray", lw=0.8, ls=":")
ax.text(pd.Timestamp("2025-08-05"), 10500, "← IS | OOS →", color="gray", fontsize=9)
ax.set_yscale("log")
ax.set_title("가상자산 $10,000 → (시간봉 과거 실시간 시뮬레이션, 수수료 0.07%, 로그 눈금)")
ax.grid(alpha=.3, which="both")
ax.legend(fontsize=9, loc="upper left")
fig.tight_layout()
fig.savefig(os.path.join(RES, "자산곡선_1000pct.png"), dpi=120)

q = eq.resample("QE").last()
qret = (q / q.shift(1).fillna(10000) - 1).mul(100).round(1)
qret.index = qret.index.to_period("Q").astype(str)

tr = pd.read_csv(os.path.join(RES, "X08_robust_X2", "X2_trades.csv"), encoding="utf-8-sig")
x2 = json.load(open(os.path.join(RES, "X08_robust_X2", "X2_summary.json"), encoding="utf-8"))
conf = pd.DataFrame([(k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
                     for k, v in x2["config"].items()], columns=["설정", "값"])

out = os.path.join(RES, "1000pct_탐색_결과.xlsx")
with pd.ExcelWriter(out, engine="openpyxl") as w:
    allr.to_excel(w, sheet_name="전체라운드", index=False)
    eq.round(2).to_excel(w, sheet_name="자산곡선")
    qret.to_excel(w, sheet_name="분기수익률(%)")
    conf.to_excel(w, sheet_name="X2설정", index=False)
    tr.to_excel(w, sheet_name="X2거래내역", index=False)
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
wb = load_workbook(out)
ws = wb.create_sheet("차트", 0)
ws.add_image(Image(os.path.join(RES, "자산곡선_1000pct.png")), "A1")
wb.save(out)
print("saved", out)
print(qret.to_string())
