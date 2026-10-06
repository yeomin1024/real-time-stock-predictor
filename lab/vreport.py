"""목표 20000%·MDD −10% 탐색 결과 묶기 → results_x/R20K_결과.xlsx · 자산곡선_R20K.png"""
import os, sys, json, pickle
import pandas as pd, numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(os.path.dirname(HERE), "results_x")
sys.path.insert(0, HERE)
import xres
plt.rcParams["font.family"] = "Malgun Gothic"
plt.rcParams["axes.unicode_minus"] = False
FINAL = ("V05_robust_P3", "P3")

KEY = [("W17_robust_P2", "P2", "P2 (MDD −5% 단계)"),
       ("V02_k_plus_rot", "K따라 k1.3", "K 따라가기 ×1.3"),
       ("V01_rot_kfilter", "R1(기준)", "6-1 모멘텀 1종목(R1_5000)"),
       ("V05_robust_P3", "P3", "P3 K×1.2 + 모멘텀 2종 — 최종"),
       ("V05_robust_P3", "P3 SP100(2023)", "P3 · 2023년 S&P100(편향 점검)")]

curves = {}
for rnd, name, label in KEY:
    e = pd.read_csv(os.path.join(RES, rnd, f"{name}_equity.csv"), encoding="utf-8-sig")
    curves[label] = e.assign(날짜=pd.to_datetime(e["날짜"])).set_index("날짜")["자산($)"]
eq = pd.DataFrame(curves)
spy = pd.read_pickle(os.path.join(HERE, "data", "spy_daily.pkl"))
prev = spy[spy.index < eq.index[0]].iloc[-1]
eq["SPY 단순보유"] = spy.reindex(eq.index).ffill() / prev * 10000
bars = xres.bars_all()
dc = pd.DataFrame({c: bars[c].groupby(bars[c]["ts"].dt.normalize())["close"].last() for c in xres.BASE58 if c in bars})
dc = dc.reindex(eq.index).ffill()
start_ok = [c for c in dc if pd.notna(dc[c].iloc[0])]
eq["58종 동일비중 단순보유"] = (dc[start_ok] / dc[start_ok].iloc[0]).mean(axis=1) * 10000

fig, ax = plt.subplots(figsize=(12, 6))
for c in eq.columns:
    ax.plot(eq.index, eq[c], label=c, lw=2.4 if "최종" in c else 1.3,
            ls="--" if "S&P100" in c else ("-." if "단순보유" in c else "-"))
ax.axhline(2010000, color="red", lw=0.8, ls=":")
ax.text(pd.Timestamp("2024-11-01"), 2150000, "목표 +20000% ($2,010,000)", color="red", fontsize=9)
ax.axvline(pd.Timestamp("2025-08-01"), color="gray", lw=0.8, ls=":")
ax.text(pd.Timestamp("2025-08-05"), 10500, "← IS | OOS →", color="gray", fontsize=9)
ax.set_yscale("log")
ax.set_title("목표 +20000% · MDD −10% — 가상자산 $10,000 → (레버리지 없음, 최신 신호, 수수료 0.07%, 로그 눈금)")
ax.grid(alpha=.3, which="both")
ax.legend(fontsize=9, loc="upper left")
fig.tight_layout()
fig.savefig(os.path.join(RES, "자산곡선_R20K.png"), dpi=120)

q = eq.resample("QE").last()
qret = (q / q.shift(1).fillna(10000) - 1).mul(100).round(1)
qret.index = qret.index.to_period("Q").astype(str)

allr = pd.read_csv(os.path.join(RES, "all_x_rounds.csv"), encoding="utf-8-sig")
allr = allr[allr["라운드"].str.startswith("V")]
cols = ["라운드", "변형", "종목수", "ALL_수익률(%)", "ALL_연환산(%)", "ALL_MDD(%)", "ALL_샤프", "IS_수익률(%)", "IS_연환산(%)",
        "IS_MDD(%)", "IS_샤프", "OOS_수익률(%)", "OOS_연환산(%)", "OOS_MDD(%)", "OOS_샤프", "ALL_청산횟수", "ALL_승률(%)",
        "ALL_손익비", "평균투자비중", "ALL_SPY수익률(%)", "판정"]
allr = allr[[c for c in cols if c in allr.columns]]
tr = pd.read_csv(os.path.join(RES, FINAL[0], f"{FINAL[1]}_trades.csv"), encoding="utf-8-sig")
summ = json.load(open(os.path.join(RES, FINAL[0], f"{FINAL[1]}_summary.json"), encoding="utf-8"))
conf = pd.DataFrame([(k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
                     for k, v in summ["config"].items()], columns=["설정", "값"])
sells = tr[tr["구분"] == "매도"].copy()
by_code = sells.groupby("종목")["손익($)"].sum().sort_values(ascending=False).round(0)
# 매도마다 진입 방식: 같은 종목의 직전 매수 사유
kind, last = [], {}
for r in tr.itertuples():
    if r.구분 == "매수":
        last[r.종목] = "K몫" if str(r.사유).startswith("K몫") else "로테이션" if str(r.사유).startswith("로테이션") else "기타"
    else:
        kind.append(last.get(r.종목, "?"))
sells["방식"] = kind
by_kind = sells.groupby("방식")["손익($)"].agg(["count", "sum", "mean"]).round(1)

out = os.path.join(RES, "R20K_결과.xlsx")
with pd.ExcelWriter(out, engine="openpyxl") as w:
    allr.to_excel(w, sheet_name="V라운드(20000%)", index=False)
    eq.round(2).to_excel(w, sheet_name="자산곡선")
    qret.to_excel(w, sheet_name="분기수익률(%)")
    by_code.to_frame("실현손익($)").to_excel(w, sheet_name="P3 종목별손익")
    by_kind.to_excel(w, sheet_name="P3 방식별손익")
    conf.to_excel(w, sheet_name="P3설정", index=False)
    tr.to_excel(w, sheet_name="P3거래내역", index=False)
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
wb = load_workbook(out)
wb.create_sheet("차트", 0).add_image(Image(os.path.join(RES, "자산곡선_R20K.png")), "A1")
wb.save(out)
print("saved", out)
print(qret.to_string())
print("final row:", {k: eq[k].iloc[-1].round(0) for k in eq})
print(by_kind.to_string())
print(by_code.head(10).to_string())
print("top5 share of P3 realized pnl:", round(by_code.head(5).sum() / by_code.sum() * 100, 1), "%")
